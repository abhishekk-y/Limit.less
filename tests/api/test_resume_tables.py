import io

from app.runtime.journey import read_document
from docx import Document


def test_docx_resume_tables_are_not_lost():
    doc = Document()
    doc.add_paragraph("Candidate experience and education")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Python SQL engineer at Example Studio"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Docker projects"
    stream = io.BytesIO()
    doc.save(stream)
    extracted = read_document(stream.getvalue(), "resume.docx")
    assert "Python SQL engineer at Example Studio" in extracted
    assert "Docker projects" in extracted
