"""Create a concise, editable judge-facing presentation from reproduced evidence."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "figures"
OUT = ROOT / "Pookie_Blinders_Team_117_Round_2_Presentation.pptx"

NAVY = "17233A"
INK = "1C2430"
MUTED = "5F6E7E"
PAPER = "FBFAF7"
WHITE = "FFFFFF"
GREEN = "A7E968"
BLUE = "82B7E3"
PURPLE = "C5A3ED"
CORAL = "E99883"
GRID = "D9DEE6"
PALE = "F3F6FA"


def rgb(hexcolor):
    h = hexcolor.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.05, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return shape


def rect(slide, x, y, w, h, fill, radius=True, line=None, width=1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    return shape


def add_bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def header(slide, kicker, title, sub=None, dark=False):
    tcolor = WHITE if dark else INK
    mcolor = "BDC9D8" if dark else MUTED
    textbox(slide, 0.62, 0.28, 12, 0.28, kicker.upper(), 10, GREEN if dark else "5C6E83", True)
    textbox(slide, 0.62, 0.62, 12.1, 0.65, title, 28, tcolor, True)
    if sub:
        textbox(slide, 0.64, 1.31, 12.0, 0.46, sub, 13, mcolor)


def footer(slide, n, dark=False):
    col = "92A0B2" if dark else "8290A0"
    textbox(slide, 0.62, 7.18, 10.4, 0.18, "LIMIT.LESS  ·  POOKIE BLINDERS  ·  TEAM 117", 8, col, True)
    textbox(slide, 12.0, 7.16, 0.7, 0.2, f"{n:02d}", 9, col, True, align=PP_ALIGN.RIGHT)


def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY if dark else PAPER)
    return slide


def bullet(slide, x, y, w, text, color=INK, dot=GREEN, size=15, h=0.62):
    rect(slide, x, y+0.11, 0.12, 0.12, dot, radius=True)
    textbox(slide, x+0.24, y, w-0.24, h, text, size, color)


def card(slide, x, y, w, h, title, body, accent=BLUE, dark=False, title_size=16, body_size=12):
    fill = "24334C" if dark else WHITE
    stroke = "34445F" if dark else GRID
    rect(slide, x, y, w, h, fill, line=stroke)
    rect(slide, x, y, 0.08, h, accent, radius=False)
    textbox(slide, x+0.24, y+0.18, w-0.46, 0.38, title, title_size, WHITE if dark else INK, True)
    textbox(slide, x+0.24, y+0.66, w-0.46, h-0.78, body, body_size, "D2DBE7" if dark else MUTED)


def metric(slide, x, y, w, value, label, accent=BLUE, dark=False, note=None):
    fill = "24334C" if dark else WHITE
    rect(slide, x, y, w, 1.16, fill, line="34445F" if dark else GRID)
    textbox(slide, x+0.18, y+0.12, w-0.36, 0.48, value, 25, accent, True)
    textbox(slide, x+0.18, y+0.61, w-0.36, 0.27, label, 10.5, WHITE if dark else INK, True)
    if note:
        textbox(slide, x+0.18, y+0.91, w-0.36, 0.18, note, 8.5, "BDC9D8" if dark else MUTED)


def add_image(slide, filename, x, y, w, h):
    path = FIG / filename
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def table(slide, x, y, w, h, headers, rows, widths=None, font_size=11, dark=False):
    shp = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tb = shp.table
    if widths:
        for col, cw in zip(tb.columns, widths):
            col.width = Inches(cw)
    for r in range(len(rows)+1):
        for c in range(len(headers)):
            cell = tb.cell(r,c)
            cell.margin_left=Inches(0.09); cell.margin_right=Inches(0.07)
            cell.margin_top=Inches(0.04); cell.margin_bottom=Inches(0.04)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r==0:
                cell.fill.fore_color.rgb=rgb(NAVY if not dark else "30415E")
                txt=headers[c]
                fg=WHITE
                bold=True
            else:
                cell.fill.fore_color.rgb=rgb(WHITE if r%2 else PALE)
                txt=rows[r-1][c]
                fg=INK
                bold=False
            tf=cell.text_frame;tf.clear();tf.word_wrap=True
            p=tf.paragraphs[0];p.space_after=Pt(0)
            run=p.add_run();run.text=str(txt);run.font.name="Aptos";run.font.size=Pt(font_size);run.font.bold=bold;run.font.color.rgb=rgb(fg)
    return tb


def build():
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    prs.core_properties.title="Limit.less Round 2 Dataset Analysis and Evidence"
    prs.core_properties.subject="SAS Data Analytics Hackathon · Team Pookie Blinders · Team 117"
    prs.core_properties.author="Pookie Blinders"
    prs.core_properties.keywords="Dataset analysis, evidence, SAS, Limit.less"

    # 1 Cover
    s=new_slide(prs,True)
    rect(s,9.4,0,3.93,7.5,"202F46",radius=False)
    rect(s,9.82,0.8,2.85,1.22,GREEN)
    textbox(s,10.0,1.03,2.5,0.55,"4 DATA LANES",19,NAVY,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    textbox(s,0.82,0.72,7.9,0.34,"SAS DATA ANALYTICS HACKATHON  ·  ROUND 2",11,GREEN,True)
    textbox(s,0.82,1.52,8.2,1.55,"Limit.less\nDataset Analysis & Evidence",34,WHITE,True)
    textbox(s,0.84,3.42,7.7,0.72,"From job-sample signals to defensible training investigations",19,"D2DBE7")
    rect(s,0.85,4.55,7.65,0.02,"526079",radius=False)
    textbox(s,0.84,4.82,7.3,0.4,"POOKIE BLINDERS  ·  TEAM ID 117",15,WHITE,True)
    textbox(s,0.84,5.35,7.7,0.48,"A reproducible analysis of four independent challenge datasets",13,"BDC9D8")
    add_image(s,"four-evidence-lanes.png",9.72,2.5,3.25,2.8)
    textbox(s,9.87,5.7,2.95,0.55,"Counts are evidence.\nClaims need boundaries.",14,WHITE,True,align=PP_ALIGN.CENTER)
    footer(s,1,True)

    # 2 The decision
    s=new_slide(prs)
    header(s,"01 · Problem & decision","What should data-career training prioritize?","The answer must survive missing fields, repeated records, uncertain skill parsing and small samples.")
    card(s,0.75,2.1,3.8,2.25,"Stakeholder decision","Help a university placement or training team choose which aggregate skill areas to investigate in a small, measurable learning pilot.",GREEN)
    card(s,4.78,2.1,3.8,2.25,"Evidence challenge","Job postings describe requested skills; the trait files describe different samples. They are not linked people, jobs or outcomes.",BLUE)
    card(s,8.8,2.1,3.8,2.25,"Research question","Which role and skill patterns are sufficiently stable to guide a training investigation—and what needs more evidence first?",PURPLE)
    textbox(s,0.85,5.0,11.7,0.75,"We do not predict who gets hired, promise salary gains, or use personality traits to screen individuals.",20,NAVY,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    footer(s,2)

    # 3 Evidence lanes
    s=new_slide(prs)
    header(s,"02 · Data design","Four datasets. Four analytical lanes.","Keep observation units separate; only compare at an explicitly reviewed concept level.")
    add_image(s,"four-evidence-lanes.png",1.2,1.95,10.95,4.8)
    footer(s,3)

    # 4 profile
    s=new_slide(prs)
    header(s,"03 · Data exploration","The dataset sizes are unequal—and the missingness matters","Descriptive counts were recalculated from the supplied CSV and Excel files.")
    add_image(s,"reproduced-source-profile.png",0.9,1.9,8.15,4.95)
    metric(s,9.45,2.05,2.85,"15,841","Analytics Jobs rows",GREEN)
    metric(s,9.45,3.38,2.85,"1,602","DataScience rows",BLUE)
    metric(s,9.45,4.71,2.85,"139 / 161","JDS / SDS observations",PURPLE)
    textbox(s,9.48,6.1,2.9,0.45,"13,806 skill cells contain `...`\nPossible truncation needs review.",10,CORAL,True)
    footer(s,4)

    # 5 market results
    s=new_slide(prs)
    header(s,"04 · Market sample","What the supplied postings actually say","Rates use exact comma-separated tokens in the available `key_skills` field—not extracted job descriptions.")
    add_image(s,"reproduced-analysis-signals.png",0.78,1.85,8.0,5.1)
    card(s,9.05,2.0,3.5,1.25,"SQL · 5.78%","915 / 15,840 available skill rows",BLUE,title_size=14,body_size=11)
    card(s,9.05,3.5,3.5,1.25,"Python · 5.30%","840 / 15,840 available skill rows",PURPLE,title_size=14,body_size=11)
    card(s,9.05,5.0,3.5,1.35,"Bengaluru · 23.7%","3,753 / 15,841 rows use first listed location",GREEN,title_size=14,body_size=10.5)
    footer(s,5)

    # 6 weighted data
    s=new_slide(prs,True)
    header(s,"05 · Weighting","93,005 is a supplied field sum—not a verified vacancy count","DataScience Jobs has no skill field; do not multiply its weights into the separate Analytics Jobs sample.",True)
    metric(s,0.88,2.15,2.8,"1,602","source rows",BLUE,True)
    metric(s,3.95,2.15,2.8,"1,460","unique references",GREEN,True)
    metric(s,7.02,2.15,2.8,"93,005","Σ num_of_jobs",PURPLE,True)
    metric(s,10.08,2.15,2.8,"3–4,200","row range · median 22",CORAL,True)
    card(s,1.0,4.05,5.3,1.7,"Company summary","TCS supplied-weight sum: 9,064 (9.7% of 93,005). Field meaning remains to be verified.",BLUE,True)
    card(s,7.0,4.05,5.3,1.7,"Title summary","Business Analyst supplied-weight sum: 32,843. It is not a count of validated distinct vacancies.",PURPLE,True)
    textbox(s,1.1,6.25,11.2,0.45,"142 excess rows have repeated reference numbers. They are retained until identifier semantics are confirmed.",14,"D2DBE7",True,align=PP_ALIGN.CENTER)
    footer(s,6,True)

    # 7 calculations
    s=new_slide(prs)
    header(s,"06 · Mathematical method","Show the arithmetic, denominator and unit","A score is only interpretable when the question and observation unit are explicit.")
    card(s,0.8,2.0,5.65,1.5,"Skill mention rate","pₖg = Σᵢ xᵢₖ / |E_g|\nRows with skill k ÷ rows with usable skill text in group g.",GREEN,title_size=15,body_size=13)
    card(s,6.85,2.0,5.65,1.5,"Supplied-weight total","W_h = Σᵢ num_of_jobsᵢ\nGrouped by company/title; shown separately from source-row count.",BLUE,title_size=15,body_size=13)
    card(s,0.8,3.85,5.65,1.6,"Trait group difference","Δₖ = mean(zₖ | y=1) − mean(zₖ | y=0)\n`r` = point-biserial correlation; Cohen's `d` = Δ / sₚ (pooled SD).",PURPLE,title_size=15,body_size=11.5)
    card(s,6.85,3.85,5.65,1.6,"SQL example","915 ÷ 15,840 = 0.0578 = 5.78%\n95% Wilson interval ≈ 5.42–6.15% under independent rows; not a national estimate.",CORAL,title_size=15,body_size=11.5)
    textbox(s,0.95,5.95,11.5,0.62,"A title/company match across files is not a valid join. No shared job, employer, person or time key is documented.",14,NAVY,True,align=PP_ALIGN.CENTER)
    footer(s,7)

    # 8 trait findings
    s=new_slide(prs)
    header(s,"07 · Separate trait research","Associations can generate questions—not causal advice","Group means were recomputed; labels and repeated IDs require careful interpretation.")
    add_image(s,"reproduced-trait-associations.png",0.75,1.8,8.4,5.15)
    card(s,9.4,2.0,3.1,1.65,"JDS signal","Largest observed mean differences: dashboard/storytelling (+1.03), maths/statistics (+0.88), coding (+0.79).",PURPLE,title_size=14,body_size=10.5)
    card(s,9.4,4.0,3.1,1.65,"SDS boundary","Aggregate personality associations only. Never feed SDS into learner, applicant or employee scoring.",BLUE,title_size=14,body_size=10.5)
    textbox(s,9.45,6.0,3.0,0.5,"No intervention effect is proven.",12,CORAL,True,align=PP_ALIGN.CENTER)
    footer(s,8)

    # 9 engine
    s=new_slide(prs,True)
    header(s,"08 · Scoring & suggestions","No opaque blended score","The engine is a transparent evidence scorecard plus decision gates—not a 0–100 employability score.",True)
    metrics=[("Demand rate","mentions ÷ eligible rows"),("Rank","percentile within source/group"),("Stability","top-k scenarios ÷ all scenarios"),("Coverage","usable text ÷ all group rows"),("Extraction quality","held-out precision / recall / F1"),("Weight influence","weighted vs row-count ranks")]
    for i,(name,desc) in enumerate(metrics):
        col=i%3;row=i//3;x=0.82+col*4.15;y=2.0+row*1.45
        card(s,x,y,3.72,1.18,name,desc,[GREEN,BLUE,PURPLE,CORAL,GREEN,BLUE][i],True,title_size=14,body_size=10.5)
    textbox(s,0.92,5.25,11.55,0.5,"Example robustness: top-10 in 8 of 10 predeclared scenarios = 80% scenario stability—not an 80% probability of importance.",14,WHITE,True,align=PP_ALIGN.CENTER)
    rect(s,0.82,5.95,11.85,0.65,"24334C",line="34445F")
    textbox(s,1.02,6.08,11.4,0.35,"SOURCE CHECK  →  REVIEW / COLLECT  →  SCENARIO GATE  →  HUMAN REVIEW  →  PILOT CANDIDATE OR DEFER",12,GREEN,True,align=PP_ALIGN.CENTER)
    footer(s,9,True)

    # 10 workflow
    s=new_slide(prs)
    header(s,"09 · Evidence workflow","A suggestion is earned through checks","If a source or measurement gate fails, the system requests evidence instead of manufacturing a recommendation.")
    add_image(s,"evidence-action-flow.png",0.95,1.82,8.45,4.95)
    card(s,9.55,2.0,2.95,1.25,"Prioritize a pilot","Only supported, reversible aggregate investigation.",GREEN,title_size=13,body_size=10.5)
    card(s,9.55,3.52,2.95,1.25,"Review mapping","A phrase, alias or duplicate choice changes the result.",PURPLE,title_size=13,body_size=10.5)
    card(s,9.55,5.04,2.95,1.25,"Collect / defer","Clarify units, IDs, coverage or validation first.",CORAL,title_size=13,body_size=10.5)
    footer(s,10)

    # 11 recommendations
    s=new_slide(prs)
    header(s,"10 · Conclusion","What we can responsibly recommend now","A descriptive signal is not yet a validated curriculum ranking.")
    bullet(s,0.9,2.05,11.5,"Investigate SQL, Python and communication/storytelling as candidate learning topics; validate skill extraction before prioritizing a module.",size=15,h=0.75)
    bullet(s,0.9,3.05,11.5,"Keep the supplied job-weight totals separate until the `num_of_jobs` definition is confirmed.",dot=BLUE,size=15,h=0.75)
    bullet(s,0.9,4.05,11.5,"Use JDS only for exploratory skill-development questions; keep SDS out of any individual recommendation or screening path.",dot=PURPLE,size=15,h=0.75)
    bullet(s,0.9,5.05,11.5,"If quality and stability gates pass, run a small voluntary pilot with baseline and follow-up skill assessments, completion and learner feedback.",dot=CORAL,size=15,h=0.85)
    textbox(s,0.95,6.35,11.5,0.35,"A short pilot can test feasibility and learning signals—not employment or salary impact.",14,NAVY,True,align=PP_ALIGN.CENTER)
    footer(s,11)

    # 12 status
    s=new_slide(prs,True)
    header(s,"11 · Evidence status","What is established—and what still needs a run receipt","The presentation does not claim the proposed SAS integration or earlier model metrics have been validated.",True)
    rows=[
        ("Source counts and descriptive summaries","Recomputed from supplied CSV/XLSX files","Available"),
        ("Exact-token rates and trait associations","Formulas, units and denominators documented","Available"),
        ("SAS / VFL execution","Programs exist; no run receipt in this report","Not verified"),
        ("Trait-model AUC / robustness","Earlier reported values lack reproducible fold/config receipt","Not independently reproduced"),
        ("Skill extraction precision / recall","Needs reviewed held-out labels; `...` export marker unresolved","Pending"),
        ("Training or employment impact","No participant pilot or outcome follow-up yet","Not measured"),
    ]
    table(s,0.85,1.95,11.65,4.55,["Evidence area","Current evidence","Status"],rows,[3.6,6.4,1.65],font_size=11,dark=True)
    textbox(s,0.95,6.65,11.35,0.28,"Next: reconcile in VFL, review the skill sample, register scenarios, and preserve run receipts.",12,GREEN,True,align=PP_ALIGN.CENTER)
    footer(s,12,True)

    # 13 close / references
    s=new_slide(prs)
    header(s,"12 · References & close","A careful dataset analysis can still drive a useful decision","Limit.less turns mixed-quality evidence into transparent, reviewable training investigations.")
    card(s,0.85,2.0,5.5,3.9,"Research basis","• SkillSpan — skill extraction benchmark\n• Cedefop Skills-OVATE — existing skills intelligence\n• Sentence-BERT — semantic candidate retrieval\n• Cawley & Talbot — model-selection bias\n• OECD — representativeness of online postings\n• Specification-curve analysis — assumption sensitivity\n\nFull citations and URLs are in the accompanying Round 2 report.",BLUE,title_size=17,body_size=12)
    card(s,6.95,2.0,5.5,3.9,"The decision we support","Use the source sample to identify questions worth testing.\n\nShow the denominator.\n\nStress-test assumptions.\n\nKeep independent datasets separate.\n\nRecommend a pilot only after evidence gates pass.",GREEN,title_size=17,body_size=15)
    textbox(s,1.0,6.3,11.3,0.5,"Pookie Blinders  ·  Team 117  ·  Limit.less",18,NAVY,True,align=PP_ALIGN.CENTER)
    footer(s,13)

    prs.save(OUT)
    print(OUT)


if __name__=="__main__":
    build()
