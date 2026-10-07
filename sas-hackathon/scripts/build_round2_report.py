from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Limitless_Research_and_Build_Report.md"
OUTPUT = ROOT / "Pookie_Blinders_Team_117_Round_2_Report.docx"


def set_cell_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9DEE6")
        borders.append(element)
    tbl_pr.append(borders)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Limit.less · Pookie Blinders · Team 117     ").font.size = Pt(8)
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_inline(paragraph, text, base_size=11.5):
    # Small, deterministic Markdown subset: bold, inline code, and emphasis.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if not (part.startswith("`") and part.endswith("`")):
            run.font.size = Pt(base_size)


def m_el(tag, attrs=None):
    node = OxmlElement(f"m:{tag}")
    for key, value in (attrs or {}).items():
        node.set(qn(f"m:{key}"), value)
    return node


def m_text(text):
    run = m_el("r")
    t = m_el("t")
    t.text = text
    run.append(t)
    return run


def m_sub(base, sub):
    node = m_el("sSub")
    node.append(m_el("sSubPr"))
    e = m_el("e"); e.append(m_text(base)); node.append(e)
    s = m_el("sub"); s.append(m_text(sub)); node.append(s)
    return node


def m_sup(base, sup):
    node = m_el("sSup")
    node.append(m_el("sSupPr"))
    e = m_el("e"); e.append(m_text(base)); node.append(e)
    s = m_el("sup"); s.append(m_text(sup)); node.append(s)
    return node


def m_frac(num, den):
    node = m_el("f")
    node.append(m_el("fPr"))
    n = m_el("num")
    for item in num if isinstance(num, list) else [m_text(num)]: n.append(item)
    d = m_el("den")
    for item in den if isinstance(den, list) else [m_text(den)]: d.append(item)
    node.append(n); node.append(d)
    return node


def m_rad(items):
    node = m_el("rad")
    props = m_el("radPr"); props.append(m_el("degHide", {"val":"1"})); node.append(props)
    node.append(m_el("deg"))
    e = m_el("e")
    for item in items if isinstance(items, list) else [m_text(items)]: e.append(item)
    node.append(e)
    return node


def m_sum(subscript, body, superscript=None):
    node = m_el("nary")
    props = m_el("naryPr")
    props.append(m_el("chr", {"val":"∑"}))
    props.append(m_el("limLoc", {"val":"undOvr"}))
    props.append(m_el("grow", {"val":"1"}))
    props.append(m_el("subHide", {"val":"0"}))
    props.append(m_el("supHide", {"val":"0" if superscript is not None else "1"}))
    node.append(props)
    sub = m_el("sub"); sub.append(m_text(subscript)); node.append(sub)
    sup = m_el("sup"); sup.append(m_el("e"))
    if superscript is not None: sup.append(m_text(superscript))
    node.append(sup)
    e = m_el("e")
    for item in body if isinstance(body, list) else [m_text(body)]: e.append(item)
    node.append(e)
    return node


def equation_items(key):
    M = m_text
    if key == "skill_rate":
        return [m_sub("p", "kg"), M(" = "), m_frac([m_sum("i = 1", [m_sub("x", "ik")], "|E_g|")], [M("|E_g|")])]
    if key == "wilson":
        # Wilson interval: a display equation with nested fraction and radical.
        phat = M("p̂")
        first = [phat, M(" + "), m_frac([m_sup("z", "2")], [M("2n")])]
        radical = m_rad([m_frac([M("p̂(1 − p̂)")], [M("n")]), M(" + "), m_frac([m_sup("z", "2")], [M("4n²")])])
        numerator = first + [M(" ± z"), radical]
        denominator = [M("1 + "), m_frac([m_sup("z", "2")], [M("n")])]
        return [M("CI₉₅% = "), m_frac(numerator, denominator)]
    if key == "weight_and_count":
        return [m_sub("n", "h"), M(" = "), m_sum("i ∈ h", [M("1")]), M("     "), m_sub("W", "h"), M(" = "), m_sum("i ∈ h", [m_sub("w", "i")])]
    if key == "trait_difference":
        return [m_sub("Δ", "k"), M(" = z̄₁ₖ − z̄₀ₖ")]
    if key == "point_biserial":
        return [m_sub("r", "pb"), M(" = "), m_frac([M("M₁ − M₀")], [M("s_z")]), m_rad([m_frac([M("n₁n₀")], [M("n²")])])]
    if key == "cohen_d":
        return [M("d = "), m_frac([m_sub("Δ", "k")], [m_sub("s", "p")])]
    if key == "pooled_sd":
        inside = m_frac([m_text("(n₁ − 1)s₁² + (n₀ − 1)s₀²")], [m_text("n₁ + n₀ − 2")])
        return [m_sub("s", "p"), M(" = "), m_rad([inside])]
    if key == "engine_demand":
        return [M("M1: Dₖg = 100 · "), m_frac([m_sub("m", "kg")], [M("|E_g|")])]
    if key == "engine_rank":
        return [M("M2: Rₖg = 100 · PercentileRank({pⱼg : j ∈ K})")]
    if key == "engine_stability":
        indicator = M("1[rankₛ(k) ≤ K]")
        return [M("M3: Sₖ = 100 · "), m_frac([m_sum("s ∈ S", [indicator])], [M("|S|")])]
    if key == "engine_coverage":
        return [M("M4: Cg = 100 · "), m_frac([M("U_g")], [M("N_g")])]
    if key == "engine_extraction":
        return [M("M5: Precision = "), m_frac([M("TP")], [M("TP + FP")]), M(";   Recall = "), m_frac([M("TP")], [M("TP + FN")]), M(";   F₁ = "), m_frac([M("2 · Precision · Recall")], [M("Precision + Recall")])]
    if key == "engine_weight":
        return [M("M6: "), m_sub("W", "h"), M(" = "), m_sum("i ∈ h", [m_sub("w", "i")]),
                M("; "), m_sub("n", "h"), M(" = |h|; "), m_sub("W", "h,−max"), M(" = "),
                m_sub("W", "h"), M(" − max"), m_sub("i", "i ∈ h"), M(" wᵢ")]
    if key == "sql_rate_example":
        return [M("p̂ = "), m_frac([M("915")], [M("15,840")]), M(" = 0.0578 = 5.78%")]
    if key == "weight_concentration":
        numerator = m_sup("(Σ_i w_i)", "2")
        denominator = M("Σ_i w_i²")
        return [M("C_w = "), m_frac([numerator], [denominator])]
    if key == "jaccard":
        return [M("J(A, B) = "), m_frac([M("|A ∩ B|")], [M("|A ∪ B|")])]
    return [M(key)]


def add_equation(doc, key):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    math_para = m_el("oMathPara")
    para_pr = m_el("oMathParaPr")
    para_pr.append(m_el("jc", {"val":"center"}))
    math_para.append(para_pr)
    math = m_el("oMath")
    for item in equation_items(key): math.append(item)
    math_para.append(math)
    p._p.append(math_para)


def add_table_equation(cell, key):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.clear_content()
    math = m_el("oMath")
    for item in equation_items(key): math.append(item)
    p._p.append(math)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in row):
            rows.append(row)
        i += 1
    return rows, i


def add_markdown_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = True
    table.style = "Table Grid"
    for ridx, row_data in enumerate(rows):
        cells = table.rows[0].cells if ridx == 0 else table.add_row().cells
        for cidx in range(cols):
            cell = cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            value = row_data[cidx] if cidx < len(row_data) else ""
            add_inline(p, value, base_size=9)
            if ridx == 0:
                set_cell_shading(cell, "233B5D")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif ridx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
        if ridx == 0:
            set_repeat_table_header(table.rows[0])
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, alt):
    match = re.match(r"!?\[([^\]]*)\]\(([^)]+)\)$", alt)
    if not match:
        return
    image_path = (SOURCE.parent / match.group(2)).resolve()
    if not image_path.exists():
        p = doc.add_paragraph(f"[Figure unavailable: {image_path.name}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(6.7))
    caption = match.group(1).strip()
    if caption:
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        for run in cp.runs:
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(80, 91, 106)


def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor(28, 35, 45)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    add_page_field(sec.footer.paragraphs[0])


def main():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure(doc)

    title = next((line[2:].strip() for line in lines if line.startswith("# ")), "Limit.less Round 2 Dataset Analysis and Evidence Report")
    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(100)
    title_p.paragraph_format.space_after = Pt(18)
    title_p.add_run(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    r = subtitle.add_run("SAS Data Analytics Hackathon · Round 2")
    r.bold = True
    r.font.size = Pt(15)
    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(team, "**Team:** Pookie Blinders  |  **Team ID:** 117", 12)
    project = doc.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(project, "**Project:** Limit.less", 12)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in lines:
        if line.startswith("**Prepared:**"):
            add_inline(date, line.replace("**", ""), 10)
            break
    doc.add_page_break()

    i = 0
    content_started = False
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("# ") or line.startswith("**Team:") or line.startswith("**SAS Data Analytics") or line.startswith("**Project:") or line.startswith("**Prepared:"):
            i += 1
            continue
        if line.startswith("!["):
            add_image(doc, line[1:])
            content_started = True
            i += 1
            continue
        equation_match = re.fullmatch(r"\[\[EQ:([a-z_]+)\]\]", line)
        if equation_match:
            add_equation(doc, equation_match.group(1))
            content_started = True
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            content_started = True
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            content_started = True
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    props = doc.core_properties
    props.title = title
    props.subject = "Round 2 dataset analysis, mathematical methodology, and evidence-based recommendations"
    props.author = "Pookie Blinders, Team 117"
    props.keywords = "Limit.less, SAS Data Analytics Hackathon, dataset analysis"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
